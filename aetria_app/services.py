"""Phase 1 domain: masters, assignment, Excel distribution, workshops."""
from __future__ import annotations

import csv
import io
import re
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from aetria_app.db import get_db

UPLOAD_ROOT = Path(__file__).resolve().parent.parent / "data" / "uploads"

# Company download Data sheet. Names match the reference workbook.
# The first five are filled from the assigned APQC master. The rest stay empty.
# Headers through Decision Point are green. Decision Criteria onward are yellow.
EXCEL_HEADERS = [
    "Level",
    "Hierarchy ID",
    "APQC ID",
    "Process / activity / task",
    "Parent ID",
    "Step ID #",
    "Task",
    "Activity Description",
    "Remarks",
    "Actor / Role",
    "Department",
    "Location",
    "Document Name",
    "Document Type",
    "System Used",
    "Transaction Code",
    "Mode",
    "Communication Mode",
    "Cycle Time",
    "Waiting Time",
    "Target Time",
    "Volume",
    "Batch or Real-Time",
    "Decision Point",
    "Decision Criteria",
    "Risk / Issue",
    "Rework",
    "Control Point",
    "KPI Impact",
    "Improvement Opportunity",
    "Automation Potential",
]
FILLED_COLUMN_COUNT = 5
GREEN_HEADER_COUNT = 24

# Reference sheet: same column names, with the workshop-field guide.
# Last flag marks the yellow columns on the Data sheet.
GUIDELINE_ROWS = [
    ("Level", "APQC level label, from L1 Category through L5 Task.", "L1 Category", False),
    ("Hierarchy ID", "Hierarchy code from the APQC master.", "4.1.1", False),
    ("APQC ID", "Published APQC element id.", "10221", False),
    ("Process / activity / task", "Element name from the APQC master.", "Develop production and materials strategies", False),
    ("Parent ID", "Hierarchy code of the parent element. Blank on a top-level category.", "4.1", False),
    ("Step ID #", "Workshop step number. Left blank for the team to fill.", "", False),
    ("Task", "Task name used in the workshop when it is finer than the APQC element.", "", False),
    ("Activity Description", "What exactly happens.", "Check specs, quantity, delivery dates", False),
    ("Remarks", "Working notes for this row.", "", False),
    ("Actor / Role", "Who performs the step.", "Sales Coordinator", False),
    ("Department", "Function that owns the step.", "Sales", False),
    ("Location", "Site or region, if needed.", "Colombo", False),
    ("Document Name", "Supporting document.", "RFQ Email / RFQ Form", False),
    ("Document Type", "Form, email, or system generated.", "Email", False),
    ("System Used", "ERP, CRM, Excel, or other system.", "SAP / Outlook", False),
    ("Transaction Code", "Transaction code when the step runs in an ERP.", "VA21", False),
    ("Mode", "How the step is executed.", "Manual / Semi / Automated", False),
    ("Communication Mode", "Email, phone, or system.", "Email", False),
    ("Cycle Time", "Time to complete the step.", "30 mins", False),
    ("Waiting Time", "Delay before the step.", "2 hrs", False),
    ("Target Time", "Expected time or SLA.", "1 day", False),
    ("Volume", "Transactions handled.", "20 RFQs/day", False),
    ("Batch or Real-Time", "Whether the step runs in batch or in real time.", "Batch", False),
    ("Decision Point", "Whether this is a decision step.", "Y", False),
    ("Decision Criteria", "Basis for the decision.", "Pricing, approval limit", True),
    ("Risk / Issue", "Known problem.", "Missing RFQ details", True),
    ("Rework", "Whether rework occurs.", "Y", True),
    ("Control Point", "Validation or check.", "Pricing approval", True),
    ("KPI Impact", "What the step affects.", "Lead time", True),
    ("Improvement Opportunity", "Observed gap.", "Manual entry duplication", True),
    ("Automation Potential", "Where the step could be automated.", "", True),
]
assert [row[0] for row in GUIDELINE_ROWS] == EXCEL_HEADERS
assert GREEN_HEADER_COUNT < len(EXCEL_HEADERS)

ROLES = ("ADMIN", "AETRIA_AREA", "COMPANY_ADMIN", "AREA_EDITOR", "VIEWER")

ROLE_LABELS = {
    "ADMIN": "Admin",
    "AETRIA_AREA": "Aetria by area (later)",
    "COMPANY_ADMIN": "Company admin (later)",
    "AREA_EDITOR": "Area editor (later)",
    "VIEWER": "Viewer (later)",
}


def new_id() -> str:
    return uuid.uuid4().hex


def utcnow() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def fetchall(sql: str, params: tuple = ()) -> list[dict]:
    with get_db() as conn:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]


def fetchone(sql: str, params: tuple = ()) -> dict | None:
    with get_db() as conn:
        row = conn.execute(sql, params).fetchone()
        return dict(row) if row else None


def file_token(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9.]+", "", (value or "").strip())
    return cleaned or "NA"


def infer_level(code: str, explicit: str | None) -> int:
    if explicit not in (None, ""):
        try:
            return max(1, min(5, int(explicit)))
        except ValueError:
            pass
    if re.fullmatch(r"\d+\.0", code):
        return 1
    return max(1, min(5, code.count(".") + (0 if code.endswith(".0") else 1)))


def infer_parent_code(code: str, known: set[str]) -> str | None:
    if re.fullmatch(r"\d+\.0", code) or "." not in code:
        return None
    parts = code.split(".")
    parent = ".".join(parts[:-1])
    if parent in known:
        return parent
    dotted = f"{parts[0]}.0"
    if len(parts) == 2 and dotted in known:
        return dotted
    return None


def _header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").strip().lower())


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _level_number(label: str, code: str) -> int:
    match = re.search(r"L(\d)", label or "", re.I)
    if match:
        return max(1, min(5, int(match.group(1))))
    return infer_level(code, None)


def read_master_rows(filename: str, data: bytes) -> list[dict]:
    """Read either the Consumer Products consolidated workbook or a simple code/name file."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = wb["All Areas Consolidated"] if "All Areas Consolidated" in wb.sheetnames else wb.active
        grid = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        header_at = None
        for index, row in enumerate(grid[:25]):
            keys = [_header_key(_cell(cell)) for cell in row]
            if "hierarchyid" in keys and ("apqcid" in keys or "processactivitytask" in keys):
                header_at = index
                break
        if header_at is not None:
            return _rows_from_apqc_sheet(grid, header_at)
        if not grid:
            raise ValueError("Workbook is empty")
        headers = [_cell(cell) for cell in grid[0]]
        raw_rows = [
            {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
            for row in grid[1:]
        ]
        return _rows_from_generic(raw_rows, headers)
    if name.endswith(".csv") or name.endswith(".txt"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        return _rows_from_generic(list(reader), list(reader.fieldnames or []))
    raise ValueError("Upload a .csv or .xlsx file")


def _rows_from_apqc_sheet(grid: list[tuple], header_at: int) -> list[dict]:
    headers = [_header_key(_cell(cell)) for cell in grid[header_at]]

    def pos(*aliases: str) -> int | None:
        for alias in aliases:
            if alias in headers:
                return headers.index(alias)
        return None

    i_area = pos("area")
    i_level = pos("level")
    i_hier = pos("hierarchyid")
    i_apqc = pos("apqcid")
    i_name = pos("processactivitytask", "process")
    i_parent = pos("parentid")
    i_applies = pos("appliestocbl")
    i_notes = pos("cblownerworkshopnotes")
    i_desc = pos("description", "remark")
    i_metrics = pos("metricsavailable")
    i_element = pos("elementdescription")
    if i_hier is None or i_name is None:
        raise ValueError("Could not find Hierarchy ID and process name columns")

    parsed = []
    for row in grid[header_at + 1 :]:
        def take(index: int | None) -> str:
            if index is None or index >= len(row):
                return ""
            return _cell(row[index])

        code = take(i_hier)
        label = take(i_level)
        pname = take(i_name)
        if not code or not pname or not re.match(r"L\d", label, re.I):
            continue
        parsed.append(
            {
                "code": code,
                "name": pname,
                "description": take(i_desc),
                "parent_code": take(i_parent),
                "level": _level_number(label, code),
                "level_label": label,
                "area": take(i_area),
                "apqc_id": take(i_apqc),
                "applies_to": take(i_applies),
                "notes": take(i_notes),
                "metrics_available": take(i_metrics),
                "element_description": take(i_element),
                "applies_to": take(i_applies),
                "notes": take(i_notes),
                "description": take(i_desc),
                "metrics_available": take(i_metrics),
                "element_description": take(i_element),
            }
        )
    if not parsed:
        raise ValueError("No APQC process rows found")
    return parsed


def _rows_from_generic(raw_rows: list[dict], headers: list[str]) -> list[dict]:
    keymap = {_header_key(h): h for h in headers if h}

    def col(*aliases: str) -> str | None:
        for alias in aliases:
            if alias in keymap:
                return keymap[alias]
        return None

    code_col = col("code", "hierarchyid", "processnumber")
    name_col = col("name", "processname", "processactivitytask")
    if not code_col or not name_col:
        raise ValueError("File needs Hierarchy ID and process name, or code and name")
    desc_col = col("description", "desc")
    parent_col = col("parentcode", "parent", "parentid")
    level_col = col("level")
    area_col = col("area")
    apqc_col = col("apqcid")
    applies_col = col("appliestocbl")
    notes_col = col("cblownerworkshopnotes", "owner")
    remark_col = col("remark", "description", "desc")
    metrics_col = col("metricsavailable")
    element_col = col("elementdescription")
    parsed = []
    for raw in raw_rows:
        code = _cell(raw.get(code_col))
        pname = _cell(raw.get(name_col))
        if not code or not pname or code.lower() in {"code", "hierarchy id"}:
            continue
        label = _cell(raw.get(level_col)) if level_col else ""
        parsed.append(
            {
                "code": code,
                "name": pname,
                "description": _cell(raw.get(remark_col)) if remark_col else (_cell(raw.get(desc_col)) if desc_col else ""),
                "parent_code": _cell(raw.get(parent_col)) if parent_col else "",
                "level": _level_number(label, code) if label else infer_level(code, None),
                "level_label": label,
                "area": _cell(raw.get(area_col)) if area_col else "",
                "apqc_id": _cell(raw.get(apqc_col)) if apqc_col else "",
                "applies_to": _cell(raw.get(applies_col)) if applies_col else "",
                "notes": _cell(raw.get(notes_col)) if notes_col else "",
                "metrics_available": _cell(raw.get(metrics_col)) if metrics_col else "",
                "element_description": _cell(raw.get(element_col)) if element_col else "",
            }
        )
    if not parsed:
        raise ValueError("No process rows found")
    return parsed


def parse_upload(filename: str, data: bytes) -> list[dict]:
    return read_master_rows(filename, data)


def bulk_upsert_master(master_id: str, rows: list[dict]) -> dict:
    created = 0
    updated = 0
    with get_db() as conn:
        existing = {
            r["code"]: r["id"]
            for r in conn.execute(
                "SELECT id, code FROM process_nodes WHERE master_id = ?",
                (master_id,),
            ).fetchall()
        }
        for index, row in enumerate(rows, start=1):
            level = int(row.get("level") or infer_level(row["code"], None))
            payload = (
                row["name"],
                row.get("description") or "",
                level,
                row.get("level_label") or "",
                row.get("area") or "",
                row.get("apqc_id") or "",
                row.get("applies_to") or "",
                row.get("notes") or "",
                row.get("metrics_available") or "",
                row.get("element_description") or "",
                index,
            )
            if row["code"] in existing:
                conn.execute(
                    """
                    UPDATE process_nodes
                    SET name=?, description=?, level=?, level_label=?, area=?, apqc_id=?,
                        applies_to=?, notes=?, metrics_available=?, element_description=?, sort_order=?
                    WHERE id=?
                    """,
                    (*payload, existing[row["code"]]),
                )
                updated += 1
            else:
                node_id = new_id()
                conn.execute(
                    """
                    INSERT INTO process_nodes
                    (id, master_id, parent_id, code, name, description, level, level_label,
                     area, apqc_id, applies_to, notes, metrics_available, element_description, sort_order)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (node_id, master_id, None, row["code"], *payload),
                )
                existing[row["code"]] = node_id
                created += 1
        known = set(existing)
        for row in rows:
            parent_code = (row.get("parent_code") or "").strip() or infer_parent_code(row["code"], known)
            parent_id = existing.get(parent_code) if parent_code else None
            if parent_id == existing.get(row["code"]):
                parent_id = None
            conn.execute(
                "UPDATE process_nodes SET parent_id = ? WHERE id = ?",
                (parent_id, existing[row["code"]]),
            )
    return {"created": created, "updated": updated}


def import_master_file(master_id: str, filename: str, data: bytes) -> dict:
    return bulk_upsert_master(master_id, read_master_rows(filename, data))


def load_nodes(master_id: str) -> list[dict]:
    raw = fetchall(
        "SELECT * FROM process_nodes WHERE master_id = ? ORDER BY sort_order, code",
        (master_id,),
    )
    by_parent: dict[str | None, list[dict]] = defaultdict(list)
    for node in raw:
        by_parent[node["parent_id"]].append(node)
    ordered: list[dict] = []

    def walk(parent_id: str | None, depth: int) -> None:
        kids = sorted(by_parent.get(parent_id, []), key=lambda n: (n["sort_order"], n["code"]))
        for node in kids:
            item = dict(node)
            item["depth"] = depth
            ordered.append(item)
            walk(node["id"], depth + 1)

    walk(None, 0)
    return ordered


def suggestion_codes(value_stream_id: str) -> list[str]:
    rows = fetchall(
        """
        SELECT process_code FROM value_stream_suggestions
        WHERE value_stream_id = ?
        ORDER BY process_code
        """,
        (value_stream_id,),
    )
    return [r["process_code"] for r in rows]


def expand_codes(master_id: str, codes: list[str]) -> tuple[list[str], list[str]]:
    """Return (node ids including descendants and ancestors, codes missing from the master)."""
    nodes = load_nodes(master_id)
    by_code = {n["code"]: n for n in nodes}
    by_id = {n["id"]: n for n in nodes}
    children: dict[str | None, list[str]] = defaultdict(list)
    for node in nodes:
        children[node["parent_id"]].append(node["id"])
    keep: set[str] = set()
    missing: list[str] = []

    def add_descendants(node_id: str) -> None:
        keep.add(node_id)
        for child_id in children.get(node_id, []):
            add_descendants(child_id)

    for code in codes:
        node = by_code.get(code)
        if not node:
            missing.append(code)
            continue
        add_descendants(node["id"])
        current = node
        while current:
            keep.add(current["id"])
            parent_id = current["parent_id"]
            current = by_id.get(parent_id) if parent_id else None
    return list(keep), missing


def save_assignment(
    company_id: str,
    master_id: str,
    value_stream_id: str,
    selected_ids: list[str],
    visible_ids: list[str] | None = None,
) -> int:
    nodes = load_nodes(master_id)
    by_id = {n["id"]: n for n in nodes}
    if visible_ids is None:
        base = set(selected_ids)
    else:
        visible = set(visible_ids)
        base = (assigned_ids(company_id, value_stream_id) - visible) | set(selected_ids)
    keep: set[str] = set()
    for node_id in base:
        current = by_id.get(node_id)
        while current:
            keep.add(current["id"])
            parent_id = current["parent_id"]
            current = by_id.get(parent_id) if parent_id else None
    now = utcnow()
    with get_db() as conn:
        conn.execute(
            "DELETE FROM assignments WHERE company_id = ? AND value_stream_id = ?",
            (company_id, value_stream_id),
        )
        for node_id in keep:
            conn.execute(
                """
                INSERT INTO assignments
                (id, company_id, master_id, value_stream_id, process_node_id, created_at)
                VALUES (?,?,?,?,?,?)
                """,
                (new_id(), company_id, master_id, value_stream_id, node_id, now),
            )
        existing = conn.execute(
            "SELECT id FROM company_streams WHERE company_id = ? AND value_stream_id = ?",
            (company_id, value_stream_id),
        ).fetchone()
        status = "assigned" if keep else "suggested"
        if existing:
            conn.execute(
                "UPDATE company_streams SET master_id = ?, status = ? WHERE id = ?",
                (master_id, status, existing["id"]),
            )
        else:
            conn.execute(
                """
                INSERT INTO company_streams (id, company_id, value_stream_id, master_id, status)
                VALUES (?,?,?,?,?)
                """,
                (new_id(), company_id, value_stream_id, master_id, status),
            )
    return len(keep)


def assigned_ids(company_id: str, value_stream_id: str) -> set[str]:
    rows = fetchall(
        "SELECT process_node_id FROM assignments WHERE company_id = ? AND value_stream_id = ?",
        (company_id, value_stream_id),
    )
    return {r["process_node_id"] for r in rows}


def distribution_rows() -> list[dict]:
    return fetchall(
        """
        SELECT c.id AS company_id, c.name AS company_name, c.short_name,
               v.id AS value_stream_id, v.code AS vs_code, v.name AS vs_name,
               m.version, m.industry, m.name AS master_name,
               COUNT(a.id) AS process_count
        FROM assignments a
        JOIN companies c ON c.id = a.company_id
        JOIN value_streams v ON v.id = a.value_stream_id
        JOIN apqc_masters m ON m.id = a.master_id
        GROUP BY c.id, c.name, c.short_name, v.id, v.code, v.name, m.version, m.industry, m.name
        ORDER BY c.name, v.sort_order, v.code
        """
    )


def _pcf_path(node: dict, by_id: dict) -> str:
    chain = []
    current = node
    seen: set[str] = set()
    while current and current["id"] not in seen:
        seen.add(current["id"])
        chain.append(f"{current['code']} {current['name']}")
        parent_id = current.get("parent_id")
        current = by_id.get(parent_id) if parent_id else None
    return " > ".join(reversed(chain))


def build_workbook(company_id: str, value_stream_id: str) -> tuple[bytes, str]:
    company = fetchone("SELECT * FROM companies WHERE id = ?", (company_id,))
    stream = fetchone("SELECT * FROM value_streams WHERE id = ?", (value_stream_id,))
    if not company or not stream:
        raise ValueError("Company or value stream not found")
    assigned = fetchall(
        """
        SELECT n.*, m.version
        FROM assignments a
        JOIN process_nodes n ON n.id = a.process_node_id
        JOIN apqc_masters m ON m.id = a.master_id
        WHERE a.company_id = ? AND a.value_stream_id = ?
        """,
        (company_id, value_stream_id),
    )
    if not assigned:
        raise ValueError("Nothing assigned for this company and value stream yet")
    master_ids = {row["master_id"] for row in assigned}
    if len(master_ids) != 1:
        raise ValueError("Assignment mixes more than one master; save the assignment again")
    master_id = next(iter(master_ids))
    version = assigned[0]["version"]
    allowed = {row["id"] for row in assigned}
    all_nodes = load_nodes(master_id)
    by_id = {node["id"]: node for node in all_nodes}
    ordered = [node for node in all_nodes if node["id"] in allowed]
    column_count = len(EXCEL_HEADERS)

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    reference = wb.create_sheet("Reference")

    header_font = Font(name="Calibri", bold=True, color="000000", size=11)
    green_header = PatternFill("solid", fgColor="70AD47")
    yellow_header = PatternFill("solid", fgColor="FFFF00")
    group_fill = PatternFill("solid", fgColor="E2EFDA")
    process_fill = PatternFill("solid", fgColor="DDEBF7")
    l1_font = Font(name="Calibri", bold=True, size=11, color="000000")
    body_font = Font(name="Calibri", size=11, color="000000")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    for col, header in enumerate(EXCEL_HEADERS, start=1):
        cell = data.cell(1, col, header)
        cell.fill = green_header if col <= GREEN_HEADER_COUNT else yellow_header
        cell.font = header_font
        cell.alignment = wrap
        cell.border = thin

    for index, node in enumerate(ordered, start=1):
        level = int(node["level"])
        parent = by_id.get(node.get("parent_id") or "")
        values = [
            node.get("level_label") or f"L{level}",
            node["code"],
            node.get("apqc_id") or "",
            node["name"],
            parent["code"] if parent else "",
        ]
        excel_row = index + 1
        for col in range(1, column_count + 1):
            cell = data.cell(excel_row, col, values[col - 1] if col <= FILLED_COLUMN_COUNT else None)
            cell.border = thin
            indent = max(0, level - 1) if col == 4 else 0
            cell.alignment = Alignment(wrap_text=True, vertical="center", indent=indent)
            cell.font = l1_font if col <= FILLED_COLUMN_COUNT and level == 1 else body_font
            if col <= FILLED_COLUMN_COUNT and level <= 2:
                cell.fill = group_fill
            elif col <= FILLED_COLUMN_COUNT and level >= 3:
                cell.fill = process_fill
        data.row_dimensions[excel_row].height = 22 if level > 1 else 26

    widths = [
        18, 16, 14, 55, 14, 12, 28, 36, 24, 18, 16, 16, 22, 16, 18, 18,
        16, 22, 14, 14, 14, 14, 20, 16, 20, 16, 12, 16, 14, 26, 22,
    ]
    for col, width in enumerate(widths, start=1):
        data.column_dimensions[get_column_letter(col)].width = width
    data.row_dimensions[1].height = 36
    data.freeze_panes = "A2"
    last_col = get_column_letter(column_count)
    data.auto_filter.ref = f"A1:{last_col}{max(data.max_row, 1)}"
    data.sheet_view.showGridLines = False
    data.oddHeader.left.text = f"{company['short_name']} · {stream['code']} · {version}"

    blue = PatternFill("solid", fgColor="1F4E79")
    blue_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    yellow = PatternFill("solid", fgColor="FFFF00")
    plain = Font(name="Calibri", size=11, color="1A1A1A")
    bold = Font(name="Calibri", bold=True, size=11, color="1A1A1A")
    for col, title in enumerate(("Column Name", "Description", "Example"), start=1):
        cell = reference.cell(1, col, title)
        cell.fill = blue
        cell.font = blue_font
        cell.alignment = wrap
        cell.border = thin
    for index, (field, meaning, example, highlight) in enumerate(GUIDELINE_ROWS, start=2):
        for col, value in enumerate((field, meaning, example), start=1):
            cell = reference.cell(index, col, value)
            cell.font = bold if col == 1 else plain
            cell.alignment = wrap
            cell.border = thin
            if highlight:
                cell.fill = yellow
        reference.row_dimensions[index].height = 32
    reference.column_dimensions["A"].width = 28
    reference.column_dimensions["B"].width = 72
    reference.column_dimensions["C"].width = 48
    reference.row_dimensions[1].height = 22
    reference.freeze_panes = "A2"
    reference.auto_filter.ref = f"A1:C{len(GUIDELINE_ROWS) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"{file_token(company['short_name'])}_{file_token(stream['code'])}_{file_token(version)}.xlsx"
    return buffer.getvalue(), filename


def dashboard_stats() -> dict:
    def count(table: str) -> int:
        row = fetchone(f"SELECT COUNT(*) AS c FROM {table}")
        return int(row["c"]) if row else 0

    return {
        "masters": count("apqc_masters"),
        "companies": count("companies"),
        "streams": count("value_streams"),
        "assignments": count("assignments"),
        "workshops": count("workshops"),
        "critical": count("critical_points"),
        "automation": count("automation_points"),
        "users": count("users"),
        "processes": count("process_nodes"),
        "downloads": len(distribution_rows()),
    }


def parse_suggestion_upload(filename: str, data: bytes) -> list[dict]:
    rows = parse_table(filename, data)
    parsed = []
    for raw in rows:
        keyed = {_header_key(k): ("" if v is None else str(v).strip()) for k, v in raw.items()}
        code = keyed.get("valuestream") or keyed.get("valuestreamcode") or keyed.get("code")
        process = keyed.get("processcode") or keyed.get("process") or keyed.get("apqc")
        note = keyed.get("note") or keyed.get("notes") or ""
        if code and process and _header_key(code) not in {"valuestream", "code"}:
            parsed.append({"stream": code, "process_code": process, "note": note})
    if not parsed:
        raise ValueError("Mapping file needs columns: valueStream, processCode")
    return parsed


def parse_table(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        return list(reader)
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(c or "") for c in rows[0]]
        return [
            {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
            for row in rows[1:]
        ]
    raise ValueError("Upload a .csv or .xlsx file")
