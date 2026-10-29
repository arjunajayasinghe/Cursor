"""Export the Phase 1 Admin database to one Excel workbook.

Usage:
    python backup_export.py
    python backup_export.py --out /path/to/aetria-phase1-full-backup.xlsx

Passwords are left out. Workshop uploads are listed by file name; the files
themselves stay in data/uploads.
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from aetria_app.db import DB_PATH, connect
from aetria_app.services import ROLE_LABELS, UPLOAD_ROOT

STORE_OUT = Path("/cursor/stores/self/media/aetria-phase1-full-backup.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1C242C")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=11, color="101418")
THIN = Border(
    left=Side(style="thin", color="D5DDE4"),
    right=Side(style="thin", color="D5DDE4"),
    top=Side(style="thin", color="D5DDE4"),
    bottom=Side(style="thin", color="D5DDE4"),
)
WRAP = Alignment(wrap_text=True, vertical="center")


def default_out() -> Path:
    if STORE_OUT.parent.is_dir():
        return STORE_OUT
    return Path(__file__).resolve().parent / "data" / "aetria-phase1-full-backup.xlsx"


def query(conn, sql: str) -> list[dict]:
    return [dict(row) for row in conn.execute(sql).fetchall()]


def opens_through(levels: str | None) -> str:
    parts = [int(part) for part in str(levels or "").split(",") if part.strip().isdigit()]
    if not parts:
        return ""
    depth = max(parts)
    return "L1" if depth == 1 else f"L1 to L{depth}"


def yes_no(value) -> str:
    return "Yes" if int(value or 0) else "No"


def write_sheet(wb: Workbook, title: str, headers: list[str], records: list[dict]) -> int:
    sheet = wb.create_sheet(title)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = THIN
    for row_index, record in enumerate(records, start=2):
        for col, header in enumerate(headers, start=1):
            value = record.get(header, "")
            if value is None:
                value = ""
            cell = sheet.cell(row_index, col, value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN
        sheet.row_dimensions[row_index].height = 18
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(max(len(headers), 1))}{max(len(records) + 1, 1)}"
    for col, header in enumerate(headers, start=1):
        widest = len(header)
        for record in records[:80]:
            widest = max(widest, min(len(str(record.get(header, "") or "")), 60))
        sheet.column_dimensions[get_column_letter(col)].width = min(max(widest + 2, 12), 56)
    return len(records)


def collect(conn) -> list[tuple[str, str, list[str], list[dict]]]:
    masters = []
    for row in query(conn, "SELECT * FROM apqc_masters ORDER BY industry, version"):
        count = conn.execute(
            "SELECT COUNT(*) AS c FROM process_nodes WHERE master_id = ?",
            (row["id"],),
        ).fetchone()["c"]
        masters.append({
            "Master": row["name"],
            "Industry": row["industry"],
            "Version": row["version"],
            "Processes": int(count),
            "Notes": row["notes"],
            "Created": row["created_at"],
            "Id": row["id"],
        })

    processes = []
    for row in query(
        conn,
        """
        SELECT n.*, m.industry, m.version, m.name AS master_name, p.code AS parent_code
        FROM process_nodes n
        JOIN apqc_masters m ON m.id = n.master_id
        LEFT JOIN process_nodes p ON p.id = n.parent_id
        ORDER BY m.version, n.sort_order, n.code
        """,
    ):
        processes.append({
            "Master": row["master_name"],
            "Version": row["version"],
            "Level": row["level_label"] or f"L{row['level']}",
            "Level number": row["level"],
            "Hierarchy ID": row["code"],
            "APQC ID": row["apqc_id"],
            "Process": row["name"],
            "Parent ID": row["parent_code"] or "",
            "Area": row["area"],
            "Applies to": row["applies_to"],
            "Description": row["description"],
            "Element description": row["element_description"],
            "Metrics available": row["metrics_available"],
            "Notes": row["notes"],
            "Sort": row["sort_order"],
            "Id": row["id"],
        })

    companies = []
    for row in query(conn, "SELECT * FROM companies ORDER BY name"):
        companies.append({
            "Company": row["name"],
            "Short name": row["short_name"],
            "Industry": row["industry"],
            "Contact": row["contact_name"],
            "Contact email": row["contact_email"],
            "Notes": row["notes"],
            "Created": row["created_at"],
            "Id": row["id"],
        })

    streams = []
    for row in query(
        conn,
        """
        SELECT v.*, c.name AS company_name, c.short_name
        FROM value_streams v
        JOIN companies c ON c.id = v.company_id
        ORDER BY c.short_name, v.sort_order, v.code
        """,
    ):
        streams.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Code": row["code"],
            "Value stream": row["name"],
            "Opens through": opens_through(row["levels"]),
            "Levels": row["levels"],
            "Description": row["description"],
            "Sort": row["sort_order"],
            "Id": row["id"],
        })

    hierarchy = []
    for row in query(
        conn,
        """
        SELECT s.process_code, s.note, v.code, v.name AS stream_name, c.name AS company_name, c.short_name
        FROM value_stream_suggestions s
        JOIN value_streams v ON v.id = s.value_stream_id
        JOIN companies c ON c.id = v.company_id
        ORDER BY c.short_name, v.code, s.process_code
        """,
    ):
        hierarchy.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"],
            "Value stream": row["stream_name"],
            "Hierarchy ID": row["process_code"],
            "Note": row["note"],
        })

    links = []
    for row in query(
        conn,
        """
        SELECT cs.status, c.name AS company_name, c.short_name, v.code, v.name AS stream_name,
               m.name AS master_name, m.version
        FROM company_streams cs
        JOIN companies c ON c.id = cs.company_id
        JOIN value_streams v ON v.id = cs.value_stream_id
        LEFT JOIN apqc_masters m ON m.id = cs.master_id
        ORDER BY c.short_name, v.code
        """,
    ):
        links.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"],
            "Value stream": row["stream_name"],
            "Master": row["master_name"] or "",
            "Version": row["version"] or "",
            "Status": row["status"],
        })

    assignments = []
    for row in query(
        conn,
        """
        SELECT c.name AS company_name, c.short_name, v.code, v.name AS stream_name,
               m.name AS master_name, m.version,
               n.level, n.level_label, n.code AS hierarchy_id, n.apqc_id, n.name AS process_name,
               p.code AS parent_code, a.created_at
        FROM assignments a
        JOIN companies c ON c.id = a.company_id
        JOIN value_streams v ON v.id = a.value_stream_id
        JOIN apqc_masters m ON m.id = a.master_id
        JOIN process_nodes n ON n.id = a.process_node_id
        LEFT JOIN process_nodes p ON p.id = n.parent_id
        ORDER BY c.short_name, v.code, n.sort_order, n.code
        """,
    ):
        assignments.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"],
            "Value stream": row["stream_name"],
            "Master": row["master_name"],
            "Version": row["version"],
            "Level": row["level_label"] or f"L{row['level']}",
            "Hierarchy ID": row["hierarchy_id"],
            "APQC ID": row["apqc_id"],
            "Process": row["process_name"],
            "Parent ID": row["parent_code"] or "",
            "Assigned": row["created_at"],
        })

    users = []
    for row in query(conn, "SELECT id, email, name, role, is_active, created_at FROM users ORDER BY email"):
        users.append({
            "Name": row["name"],
            "Email": row["email"],
            "Role": ROLE_LABELS.get(row["role"], row["role"]),
            "Role code": row["role"],
            "Active": yes_no(row["is_active"]),
            "Created": row["created_at"],
            "Id": row["id"],
        })

    workshops = []
    for row in query(
        conn,
        """
        SELECT w.*, c.name AS company_name, c.short_name, v.code, v.name AS stream_name
        FROM workshops w
        JOIN companies c ON c.id = w.company_id
        LEFT JOIN value_streams v ON v.id = w.value_stream_id
        ORDER BY w.held_on, c.short_name
        """,
    ):
        workshops.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"] or "",
            "Value stream": row["stream_name"] or "",
            "Title": row["title"],
            "Held on": row["held_on"],
            "Summary": row["summary"],
            "Created": row["created_at"],
            "Id": row["id"],
        })

    files = []
    for row in query(
        conn,
        """
        SELECT f.*, w.title, c.short_name
        FROM workshop_files f
        JOIN workshops w ON w.id = f.workshop_id
        JOIN companies c ON c.id = w.company_id
        ORDER BY f.uploaded_at
        """,
    ):
        path = UPLOAD_ROOT / row["stored_name"]
        files.append({
            "Company short name": row["short_name"],
            "Workshop": row["title"],
            "Kind": row["kind"],
            "File name": row["original_name"],
            "Stored name": row["stored_name"],
            "File on disk": "Yes" if path.is_file() else "No",
            "Uploaded": row["uploaded_at"],
        })

    critical = []
    for row in query(
        conn,
        """
        SELECT cp.*, c.name AS company_name, c.short_name, v.code, w.title AS workshop_title,
               n.code AS hierarchy_id, n.name AS process_name
        FROM critical_points cp
        JOIN companies c ON c.id = cp.company_id
        LEFT JOIN value_streams v ON v.id = cp.value_stream_id
        LEFT JOIN workshops w ON w.id = cp.workshop_id
        LEFT JOIN process_nodes n ON n.id = cp.process_node_id
        ORDER BY c.short_name, cp.created_at
        """,
    ):
        critical.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"] or "",
            "Workshop": row["workshop_title"] or "",
            "Hierarchy ID": row["hierarchy_id"] or "",
            "Process": row["process_name"] or "",
            "Title": row["title"],
            "Detail": row["detail"],
            "Owner": row["owner"],
            "Status": row["status"],
            "Created": row["created_at"],
        })

    automation = []
    for row in query(
        conn,
        """
        SELECT ap.*, c.name AS company_name, c.short_name, v.code,
               n.code AS hierarchy_id, n.name AS process_name
        FROM automation_points ap
        JOIN companies c ON c.id = ap.company_id
        LEFT JOIN value_streams v ON v.id = ap.value_stream_id
        LEFT JOIN process_nodes n ON n.id = ap.process_node_id
        ORDER BY c.short_name, ap.created_at
        """,
    ):
        automation.append({
            "Company": row["company_name"],
            "Short name": row["short_name"],
            "Value stream code": row["code"] or "",
            "Hierarchy ID": row["hierarchy_id"] or "",
            "Process": row["process_name"] or "",
            "Title": row["title"],
            "Detail": row["detail"],
            "Manual or system": row["manual_vs_system"],
            "Tool": row["tool"],
            "Priority": row["priority"],
            "Status": row["status"],
            "Created": row["created_at"],
        })

    return [
        ("Masters", "APQC masters loaded in the workbench.", list(masters[0]) if masters else [
            "Master", "Industry", "Version", "Processes", "Notes", "Created", "Id"
        ], masters),
        ("Processes", "Full APQC process tree, L1 through L5.", list(processes[0]) if processes else [
            "Master", "Version", "Level", "Level number", "Hierarchy ID", "APQC ID", "Process",
            "Parent ID", "Area", "Applies to", "Description", "Element description",
            "Metrics available", "Notes", "Sort", "Id",
        ], processes),
        ("Companies", "Companies and contacts.", list(companies[0]) if companies else [
            "Company", "Short name", "Industry", "Contact", "Contact email", "Notes", "Created", "Id"
        ], companies),
        ("Value streams", "Each company's streams and how far the tree opens.", list(streams[0]) if streams else [
            "Company", "Short name", "Code", "Value stream", "Opens through", "Levels", "Description", "Sort", "Id"
        ], streams),
        ("Hierarchy IDs", "Hierarchy IDs saved on a value stream. Empty when none were added.", list(hierarchy[0]) if hierarchy else [
            "Company", "Short name", "Value stream code", "Value stream", "Hierarchy ID", "Note"
        ], hierarchy),
        ("Stream links", "Company, value stream, and master linked for distribution.", list(links[0]) if links else [
            "Company", "Short name", "Value stream code", "Value stream", "Master", "Version", "Status"
        ], links),
        ("Assignments", "Processes assigned to a company and value stream.", list(assignments[0]) if assignments else [
            "Company", "Short name", "Value stream code", "Value stream", "Master", "Version",
            "Level", "Hierarchy ID", "APQC ID", "Process", "Parent ID", "Assigned",
        ], assignments),
        ("Users", "People who can sign in. Passwords are not included.", list(users[0]) if users else [
            "Name", "Email", "Role", "Role code", "Active", "Created", "Id"
        ], users),
        ("Workshops", "Workshop titles, dates, and summaries.", list(workshops[0]) if workshops else [
            "Company", "Short name", "Value stream code", "Value stream", "Title", "Held on", "Summary", "Created", "Id"
        ], workshops),
        ("Workshop files", "Minutes and diagrams by file name. The files stay in data/uploads.", list(files[0]) if files else [
            "Company short name", "Workshop", "Kind", "File name", "Stored name", "File on disk", "Uploaded"
        ], files),
        ("Critical points", "Critical points recorded against a company.", list(critical[0]) if critical else [
            "Company", "Short name", "Value stream code", "Workshop", "Hierarchy ID", "Process",
            "Title", "Detail", "Owner", "Status", "Created",
        ], critical),
        ("Automation", "Automation points recorded against a company.", list(automation[0]) if automation else [
            "Company", "Short name", "Value stream code", "Hierarchy ID", "Process", "Title", "Detail",
            "Manual or system", "Tool", "Priority", "Status", "Created",
        ], automation),
    ]


def build(conn) -> Workbook:
    sheets = collect(conn)
    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    written = []
    for title, about, headers, records in sheets:
        count = write_sheet(wb, title, headers, records)
        written.append((title, count, about))
        notes.append((title, about))
    headers = ["Sheet", "Rows", "What it holds"]
    for col, header in enumerate(headers, start=1):
        cell = contents.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = THIN
    contents.cell(2, 1, "Contents")
    contents.cell(2, 2, "")
    contents.cell(2, 3, "This sheet. What each tab holds and how many rows were exported.")
    for cell in contents[2]:
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = THIN
    for index, (title, count, about) in enumerate(written, start=3):
        for col, value in enumerate((title, count, about), start=1):
            cell = contents.cell(index, col, value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN
    extra = len(written) + 3
    contents.cell(extra, 1, "Note").font = BODY_FONT
    contents.cell(extra, 3, "Passwords are not in this workbook. Workshop files are named here; the files stay in data/uploads.")
    for cell in contents[extra]:
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = THIN
    contents.freeze_panes = "A2"
    contents.column_dimensions["A"].width = 22
    contents.column_dimensions["B"].width = 12
    contents.column_dimensions["C"].width = 78
    contents.row_dimensions[1].height = 22
    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Export the Phase 1 Admin database to Excel.")
    parser.add_argument("--out", type=Path, default=default_out(), help="Workbook path")
    args = parser.parse_args()
    if not DB_PATH.is_file() and not os.environ.get("DATABASE_URL"):
        raise SystemExit(f"No database at {DB_PATH}. Run python seed.py first.")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    conn = connect()
    try:
        wb = build(conn)
    finally:
        conn.close()
    wb.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
